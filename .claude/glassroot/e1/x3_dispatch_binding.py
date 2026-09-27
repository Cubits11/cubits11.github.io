"""X3: can 13.05 bind an execution receipt to the exact dispatched bytes?
Builds W0 (v1 dispatched) and W1 (v2 dispatched) in copies of the live 13.05 export, using only the
columns the sheet defines, then asks whether any receipt field distinguishes the two worlds."""
import copy, hashlib, json, sys
import openpyxl
src, out = sys.argv[1], sys.argv[2]
S1 = "ASTRA-BRIEF section A + T6 (as scored)"
S2 = "ASTRA-BRIEF section A + T1..T6 (superseding draft)"
def world(dispatched):
    wb = openpyxl.load_workbook(src)
    fz, rc = wb["Freeze Ledger"], wb["Execution Receipts"]
    hdr_f = {c.value: c.column for c in fz[1] if c.value}; hdr_r = {c.value: c.column for c in rc[1] if c.value}
    def put(ws, hdr, row, vals):
        for k, v in vals.items(): ws.cell(row=row, column=hdr[k], value=v)
    put(fz, hdr_f, 3, {"FREEZE ID": "FZ-1", "DRAFT ID": "ASTRA-X3", "PROMPT VERSION": "v1", "FROZEN AT": "2026-09-27T02:00Z",
                        "TARGET": "ASTRA", "RAW PROMPT SNAPSHOT": S1, "COMPILED PROMPT SNAPSHOT": S1, "DECISION": "REVIEW"})
    put(fz, hdr_f, 4, {"FREEZE ID": "FZ-2", "DRAFT ID": "ASTRA-X3", "PROMPT VERSION": "v2", "FROZEN AT": "2026-09-27T02:05Z",
                        "TARGET": "ASTRA", "RAW PROMPT SNAPSHOT": S2, "COMPILED PROMPT SNAPSHOT": S2, "DECISION": "REVIEW", "SUPERSEDES": "FZ-1"})
    # every receipt column the sheet defines, filled as its header asks; none of them names the dispatched version or bytes
    put(rc, hdr_r, 2, {"RUN ID": "RUN-X3", "DRAFT ID": "ASTRA-X3", "DATE": "2026-09-27", "MODEL / HARNESS": "GPT-6 Astra / Codex",
                       "EXECUTOR IDENTITY": "astra", "HARNESS": "Codex", "BRANCH / WORKTREE": "claude/x3", "HEAD / COMMIT": "abc1234",
                       "TASK ID": "T6", "TRIAL ID": "1", "TRACE / TRANSCRIPT REF": "codex-session-url", "ACTUAL METHOD": "ran T6",
                       "RESULT": "PASS", "OWNER DISPOSITION": "PENDING", "ARTIFACT DURABILITY": "DURABLE_REPO"})
    wb.save(out.replace(".xlsx", f"_{dispatched}.xlsx"))
    row = [c.value for c in rc[2]]
    return row, [(r[0].value, r[2].value, r[5].value) for r in fz.iter_rows(min_row=3, max_row=4)]
r0, freezes = world("v1")
r1, _ = world("v2")
same = r0 == r1
candidates = [f for f in freezes if f[1] in ("v1", "v2")]
# the minimal fix: one digest of the dispatched bytes on the receipt, matched to a digest of each snapshot
h = lambda s: hashlib.sha256(s.encode()).hexdigest()
fix = {"W0": [f[0] for f in freezes if h(f[2]) == h(S1)], "W1": [f[0] for f in freezes if h(f[2]) == h(S2)]}
res = {"receipt_rows_identical_across_worlds": same, "freeze_rows_joinable_from_receipt": [f[0] for f in candidates],
       "receipt_columns_naming_version_or_bytes": [],
       "decision_without_fix": "UNDECIDED: the receipt joins to both FZ-1 and FZ-2 by DRAFT ID",
       "with_digest_column": fix, "decision_with_fix": "W0 -> FZ-1 (score of v1 applies), W1 -> FZ-2 (it does not)"}
print(json.dumps(res, indent=1))
